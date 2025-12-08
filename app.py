import os
import json
import streamlit as st
from sentence_transformers import SentenceTransformer
import PyPDF2
import docx
from io import BytesIO
import requests
import traceback
from dotenv import load_dotenv
import re
import html
import numpy as np
import pathlib
import shutil
import hashlib
from typing import List, Dict, Any

# chroma import
try:
    import chromadb
except Exception:
    chromadb = None

load_dotenv()

# ------------------- Configuration -------------------
# LLM SERVICE 1: OpenRouter 
OPENROUTER_API_KEY = os.getenv("OPENROUTER_API_KEY")
OPENROUTER_URL = "https://openrouter.ai/api/v1/chat/completions"
OPENROUTER_MODEL = "amazon/nova-2-lite-v1:free"

# LLM SERVICE 2: Gemini
GEMINI_API_KEY = os.getenv("GEMINI_API_KEY") 
GEMINI_URL = "https://generativelanguage.googleapis.com/v1beta/models/{}:generateContent"
GEMINI_MODEL = "gemini-2.5-flash"


# Directories for Chroma persistence and local store
CHROMA_DB_DIR = os.path.expanduser("~/.local/share/candidate_matcher/chroma_db")
LOCAL_STORE_DIR = os.path.join(os.getcwd(), "chroma_store")
CHROMA_COLLECTION = "resumes"

os.makedirs(CHROMA_DB_DIR, exist_ok=True)
os.makedirs(LOCAL_STORE_DIR, exist_ok=True)

# ------------------- Embedding model -------------------
@st.cache_resource
def load_embedder():
    return SentenceTransformer("all-mpnet-base-v2")

embedder = load_embedder()
EMBED_DIM = 768

# ------------------- Local store files -------------------
EMBED_FILE = os.path.join(LOCAL_STORE_DIR, "embeddings.npy")
META_FILE = os.path.join(LOCAL_STORE_DIR, "metadata.json")
HASH_FILE = os.path.join(LOCAL_STORE_DIR, "hash_index.json")
ID_LIST_FILE = os.path.join(LOCAL_STORE_DIR, "id_list.json")

# in-memory structures
_local_embeddings, _local_metadata = None, None  
_id_list: List[int] = []
_next_id = 0
_hash_to_id = {}

# ------------------- Chroma client placeholders -------------------
_chroma_client = None
_chroma_collection = None

# ------------------- Helpers: local store persistence -------------------

def _load_local_store():
    global _local_embeddings, _local_metadata, _id_list, _next_id, _hash_to_id
    try:
        if os.path.exists(EMBED_FILE) and os.path.exists(META_FILE) and os.path.exists(ID_LIST_FILE):
            emb = np.load(EMBED_FILE)
            with open(META_FILE, "r", encoding="utf-8") as f:
                meta = json.load(f)
            with open(ID_LIST_FILE, "r", encoding="utf-8") as f:
                id_list = json.load(f)
            if os.path.exists(HASH_FILE):
                with open(HASH_FILE, "r", encoding="utf-8") as f:
                    hash_map = json.load(f)
            else:
                hash_map = {}
            meta = {int(k): v for k, v in meta.items()}
            id_list = [int(x) for x in id_list]
            _local_embeddings = emb.astype(np.float32)
            _local_metadata = meta
            _id_list = id_list
            _hash_to_id = {k: int(v) for k, v in hash_map.items()}
            _next_id = max(_id_list) + 1 if _id_list else 0
            return
    except Exception as e:
        st.warning(f"Failed to load local store (continuing with empty store): {e}")
    _local_embeddings = np.zeros((0, EMBED_DIM), dtype=np.float32)
    _local_metadata = {}
    _id_list = []
    _next_id = 0
    _hash_to_id = {}

def _save_local_store():
    global _local_embeddings, _local_metadata, _id_list, _hash_to_id
    try:
        if _local_embeddings is None:
            arr = np.zeros((0, EMBED_DIM), dtype=np.float32)
        else:
            arr = _local_embeddings.astype(np.float32)
        np.save(EMBED_FILE, arr)
        meta_str_keys = {str(k): v for k, v in _local_metadata.items()}
        with open(META_FILE, "w", encoding="utf-8") as f:
            json.dump(meta_str_keys, f, ensure_ascii=False, indent=2)
        with open(ID_LIST_FILE, "w", encoding="utf-8") as f:
            json.dump([int(x) for x in _id_list], f)
        with open(HASH_FILE, "w", encoding="utf-8") as f:
            json.dump({k: int(v) for k, v in _hash_to_id.items()}, f)
    except Exception as e:
        st.warning(f"Failed to save local store: {e}")

# initialize local store
_load_local_store()

# ------------------- Chroma init (modern PersistentClient) -------------------

import tempfile
import errno
import shutil
import time

def init_chroma(prefers_home: bool = True):
    global _chroma_client, _chroma_collection, CHROMA_DB_DIR

    if chromadb is None:
        _chroma_client = None
        _chroma_collection = None
        st.info("Chroma not installed — running local-only.")
        return

    candidates = [
        CHROMA_DB_DIR,
        os.path.expanduser("~/.local/share/candidate_matcher/chroma_db"),
        os.path.join(tempfile.gettempdir(), "candidate_matcher_chroma_db")
    ]

    success_path = None

    for p in candidates:
        try:
            os.makedirs(p, exist_ok=True)
        except:
            continue

        try:
            testfile = os.path.join(p, ".write_test")
            with open(testfile, "w") as f:
                f.write("ok")
            os.remove(testfile)
        except:
            continue

        try:
            try:
                client = chromadb.PersistentClient(path=p)
            except:
                client = chromadb.Client()

            try:
                coll = client.get_collection(CHROMA_COLLECTION)
            except:
                coll = client.create_collection(CHROMA_COLLECTION)

            _chroma_client = client
            _chroma_collection = coll
            success_path = p
            break
        except:
            continue

    if success_path:
        pass
    else:
        st.warning("Chroma could not be initialized — falling back to local-only.")
        _chroma_client = None
        _chroma_collection = None


def reset_chroma_collection(force_delete_dir: bool = False):
    """
    Reset/clear the Chroma collection in a safe way:
    1) If chroma client exists, try to delete the collection via API (preferred).
    2) Re-create an empty collection (client.create_collection).
    3) If client doesn't exist or deletion via API fails, attempt to remove DB files
       under CHROMA_DB_DIR safely (but only when force_delete_dir=True).
    4) Always clear local numpy store files and in-memory structures.
    """
    global _chroma_client, _chroma_collection
    global _local_embeddings, _local_metadata, _id_list, _next_id, _hash_to_id

    # -------- CLEAR LOCAL NUMPY STORE --------
    try:
        for f in [EMBED_FILE, META_FILE, ID_LIST_FILE, HASH_FILE]:
            if os.path.exists(f):
                os.remove(f)
    except Exception as e:
        print("Warning while clearing local files:", e)

    _local_embeddings = np.zeros((0, EMBED_DIM), dtype=np.float32)
    _local_metadata = {}
    _id_list = []
    _next_id = 0
    _hash_to_id = {}

    # -------- CHROMA CLIENT RESET --------
    if chromadb is None:
        return

    client = _chroma_client

    if client is None:
        try:
            try:
                client = chromadb.PersistentClient(path=CHROMA_DB_DIR)
            except:
                client = chromadb.Client()
        except:
            client = None

    if client is not None:
        try:
            if hasattr(client, "delete_collection"):
                try:
                    client.delete_collection(collection_name=CHROMA_COLLECTION)
                except TypeError:
                    client.delete_collection(CHROMA_COLLECTION)
        except Exception as e:
            print("Chroma API deletion failed:", e)

        try:
            try:
                client = chromadb.PersistentClient(path=CHROMA_DB_DIR)
            except:
                client = chromadb.Client()
            _chroma_client = client
            _chroma_collection = client.create_collection(CHROMA_COLLECTION)
            return
        except Exception:
            pass  

    if force_delete_dir:
        try:
            if os.path.exists(CHROMA_DB_DIR):
                shutil.rmtree(CHROMA_DB_DIR)
            os.makedirs(CHROMA_DB_DIR, exist_ok=True)

            try:
                _chroma_client = chromadb.PersistentClient(path=CHROMA_DB_DIR)
            except:
                _chroma_client = chromadb.Client()

            _chroma_collection = _chroma_client.create_collection(CHROMA_COLLECTION)
        except Exception as e:
            print("Full Chroma dir reset failed:", e)
            _chroma_client = None
            _chroma_collection = None

init_chroma()

# ------------------- Utility functions -------------------

def extract_text_from_file(uploaded_file) -> str:
    name = uploaded_file.name.lower()
    try:
        if name.endswith(".pdf"):
            reader = PyPDF2.PdfReader(uploaded_file)
            pages = [p.extract_text() or "" for p in reader.pages]
            return "\n".join(pages)
        elif name.endswith(".docx") or name.endswith('.doc'):
            uploaded_file.seek(0)
            doc = docx.Document(BytesIO(uploaded_file.read()))
            paragraphs = [p.text for p in doc.paragraphs]
            return "\n".join(paragraphs)
        else:
            uploaded_file.seek(0)
            return uploaded_file.read().decode("utf-8", errors="ignore")
    except Exception as e:
        st.warning(f"Failed to extract text from {uploaded_file.name}: {e}")
        return ""

def embed_text_np(text: str) -> np.ndarray:
    vec = embedder.encode(text, convert_to_tensor=False)
    return np.asarray(vec, dtype=np.float32)

def _text_hash(text: str) -> str:
    return hashlib.sha256(text.encode("utf-8")).hexdigest()

# ------------------- Add / search / reset functions -------------------

def add_resume_to_store(filename: str, full_text: str) -> int:
    """Add resume to local store and to chroma (best-effort). Returns assigned id."""
    global _local_embeddings, _local_metadata, _id_list, _next_id, _hash_to_id, _chroma_collection

    h = _text_hash(full_text)
    if h in _hash_to_id:
        existing_id = _hash_to_id[h]
        return existing_id

    rid = _next_id
    _next_id += 1

    emb = embed_text_np(full_text)
    preview = full_text[:1000]

    if _local_embeddings.size == 0:
        _local_embeddings = emb.reshape(1, -1)
    else:
        _local_embeddings = np.vstack([_local_embeddings, emb.reshape(1, -1)])

    _local_metadata[rid] = {"name": filename, "preview": preview}
    _id_list.append(rid)
    _hash_to_id[h] = rid

    # persist local store
    _save_local_store()

    if _chroma_collection is not None:
        try:
            _chroma_collection.add(
                ids=[str(rid)],
                documents=[preview],
                metadatas=[{"name": filename, "hash": h}],
                embeddings=[emb.tolist()],
            )
            try:
                _chroma_client.persist()
            except Exception:
                pass
        except Exception as e:
            st.warning(f"Warning: adding to Chroma failed for {filename}: {e}")

    return rid

def cosine_sim(a: np.ndarray, B: np.ndarray) -> np.ndarray:
    if a.ndim == 1:
        a_norm = a / (np.linalg.norm(a) + 1e-12)
    else:
        a_norm = a.reshape(-1) / (np.linalg.norm(a) + 1e-12)
    B_norm = B / (np.linalg.norm(B, axis=1, keepdims=True) + 1e-12)
    sims = (B_norm @ a_norm).reshape(-1)
    return sims

def search_local_by_cosine(jd_text: str, top_k: int = 10):
    if _local_embeddings.size == 0:
        return []
    qvec = embed_text_np(jd_text)
    sims = cosine_sim(qvec, _local_embeddings)
    top_idx = np.argsort(-sims)[:top_k]
    results = []
    for pos in top_idx:
        if pos >= len(_id_list):
            continue
        rid = _id_list[pos]
        meta = _local_metadata.get(rid, {})
        results.append({"id": rid, "score": float(sims[pos]), "payload": meta, "preview": meta.get("preview", "")})
    return results

def search_via_chroma_first_then_local(jd_text: str, top_k: int = 10):
    if _chroma_collection is not None:
        try:
            qvec = embed_text_np(jd_text).tolist()
            resp = _chroma_collection.query(
                query_embeddings=[qvec],
                n_results=top_k,
                include=["documents", "metadatas", "distances"],
            )
            ids = resp.get("ids", [[]])[0] if isinstance(resp.get("ids", None), list) else []
            docs = resp.get("documents", [[]])[0] if isinstance(resp.get("documents", None), list) else []
            metadatas = resp.get("metadatas", [[]])[0] if isinstance(resp.get("metadatas", None), list) else []
            distances = resp.get("distances", [[]])[0] if isinstance(resp.get("distances", None), list) else []

            results = []
            for i, sid in enumerate(ids):
                rid = None
                try:
                    rid = int(sid)
                except Exception:
                    try:
                        rid = int(str(sid))
                    except Exception:
                        rid = None
                payload = metadatas[i] if i < len(metadatas) else {}
                doc_preview = docs[i] if i < len(docs) else payload.get("preview", "")
                sim_val = None
                if rid is not None and rid in _local_metadata:
                    pos = _id_list.index(rid)
                    emb = _local_embeddings[pos]
                    qemb = embed_text_np(jd_text)
                    sim_val = float(cosine_sim(qemb, emb.reshape(1, -1))[0])
                else:
                    if i < len(distances) and distances[i] is not None:
                        try:
                            sim_val = float(1.0 - float(distances[i]))
                        except Exception:
                            sim_val = 0.0
                    else:
                        sim_val = 0.0
                results.append({"id": rid, "score": sim_val, "payload": payload, "preview": doc_preview})
            if not results:
                return search_local_by_cosine(jd_text, top_k)
            results = sorted(results, key=lambda x: (x["score"] if x["score"] is not None else -999.0), reverse=True)
            return results[:top_k]
        except Exception as e:
            st.warning(f"Chroma query failed (falling back to local): {e}")
            return search_local_by_cosine(jd_text, top_k)
    else:
        return search_local_by_cosine(jd_text, top_k)

def dedupe_results_by_id_keep_best(results: List[Dict[str, Any]]):
    best = {}
    for r in results:
        rid = r.get("id")
        if rid is None:
            rid = r.get("payload", {}).get("name") or r.get("preview")
        score = r.get("score") if r.get("score") is not None else 0.0
        if rid not in best or score > best[rid]["score"]:
            best[rid] = {"item": r, "score": float(score)}
    deduped = [v["item"] for v in sorted(best.values(), key=lambda x: x["score"], reverse=True)]
    return deduped

def reset_local_store():
    global _local_embeddings, _local_metadata, _id_list, _next_id, _hash_to_id
    try:
        for f in [EMBED_FILE, META_FILE, ID_LIST_FILE, HASH_FILE]:
            if os.path.exists(f):
                os.remove(f)
    except Exception as e:
        st.warning(f"Failed to remove local store files: {e}")
    _local_embeddings = np.zeros((0, EMBED_DIM), dtype=np.float32)
    _local_metadata = {}
    _id_list = []
    _next_id = 0
    _hash_to_id = {}
    _save_local_store()


# ------------------- Unified LLM Helpers -------------------

import time
from functools import lru_cache

@lru_cache(maxsize=64)
def _cached_call_llm_cache_key(service, system_prompt, user_prompt):
    return None

def call_openrouter_unified(system_prompt: str, user_prompt: str, max_retries: int = 4, base_backoff: float = 1.0):
    """Robust OpenRouter caller with exponential backoff."""
    if not OPENROUTER_API_KEY:
        st.error("OpenRouter API Key is not configured.")
        return None

    headers = {
        "Authorization": f"Bearer {OPENROUTER_API_KEY}",
        "Content-Type": "application/json",
    }
    payload = {
        "model": OPENROUTER_MODEL,
        "messages": [
            {"role": "system", "content": system_prompt},
            {"role": "user", "content": user_prompt},
        ],
        "temperature": 0.0,
    }

    last_error = None
    attempt = 0
    while attempt <= max_retries:
        attempt += 1
        try:
            resp = requests.post(OPENROUTER_URL, headers=headers, json=payload, timeout=60)
            resp.raise_for_status()
            data = resp.json()
            if "choices" in data and len(data["choices"]) > 0:
                choice = data["choices"][0]
                return choice.get("message", {}).get("content") or choice.get("text")
            return json.dumps(data)
        except requests.exceptions.HTTPError as he:
            status = getattr(he.response, "status_code", None)
            last_error = he
            if status == 429 or (status and 500 <= status < 600):
                backoff = base_backoff * (2 ** (attempt - 1))
                time.sleep(min(backoff * (1 + 0.1 * (time.time() % 1)), 30.0))
                continue
            st.error(f"OpenRouter HTTP error {status}: {getattr(he.response, 'text', None)}")
            return None
        except (requests.exceptions.Timeout, requests.exceptions.ConnectionError) as conn_e:
            last_error = conn_e
            time.sleep(min(base_backoff * (2 ** (attempt - 1)), 30.0))
            continue
        except Exception as e:
            last_error = e
            break

    st.warning(f"OpenRouter is unavailable or rate-limited. (Retries exhausted.) Last error: {repr(last_error)}")
    return None

def call_gemini(system_prompt: str, user_prompt: str):
    """Simple synchronous Gemini API caller."""
    if not GEMINI_API_KEY:
        st.error("Gemini API Key is not configured.")
        return None

    url = f"https://generativelanguage.googleapis.com/v1beta/models/{GEMINI_MODEL}:generateContent"

    headers = {
        "Content-Type": "application/json",
    }
    
    payload = {
        "contents": [
            {"role": "user", "parts": [
                {"text": f"SYSTEM PROMPT: {system_prompt}\n\nUSER PROMPT: {user_prompt}"}
            ]}
        ],
        "generationConfig": { 
            "temperature": 0.0 
        }
    }
    
    try:
        resp = requests.post(f"{url}?key={GEMINI_API_KEY}", headers=headers, json=payload, timeout=60)
        resp.raise_for_status()
        data = resp.json()
        
        if 'candidates' in data and data['candidates']:
            candidate = data['candidates'][0]
            if 'content' in candidate and 'parts' in candidate['content']:
                return candidate['content']['parts'][0]['text']
        
        st.warning(f"Gemini API returned unexpected structure: {json.dumps(data)}")
        return None
        
    except requests.exceptions.HTTPError as he:
        status = getattr(he.response, "status_code", None)
        body = getattr(he.response, 'text', 'No body')
        st.error(f"Gemini HTTP error {status}: {body}")
        return None
    except Exception as e:
        st.error(f"Gemini API call failed: {e}")
        return None


def call_llm(service: str, system_prompt: str, user_prompt: str):
    """Routes the call to the selected LLM service."""
    if service == "OpenRouter":
        return call_openrouter_unified(system_prompt, user_prompt)
    elif service == "Gemini":
        return call_gemini(system_prompt, user_prompt)
    return None


def parse_jd_with_llm(jd_text: str, service: str) -> dict:
    """
    Use LLM to parse JD -> structured JSON.
    """
    safe_empty = {"must_have": [], "important": [], "nice_to_have": [], "implicit": []}
    
    if not service: return safe_empty

    system_prompt = (
        "You are a recruitment assistant. Given a job description, output STRICT JSON only. "
        "The JSON must have keys: must_have, important, nice_to_have, implicit. "
        "Each key maps to an array of objects with schema: "
        "{\"skill\": str, \"description\": str, \"weight\": float, \"min_years\": int|null}. "
        "Return only valid JSON — do NOT include any explanation or surrounding text."
    )
    user_prompt = f"JOB DESCRIPTION:\n\n{jd_text}\n\nReturn JSON only."
    raw = None
    try:
        raw = call_llm(service, system_prompt, user_prompt)
    except Exception as e:
        st.warning(f"{service} call raised an exception: {e}")
        raw = None

    if not raw:
        lines = [l.strip("-• \t") for l in jd_text.splitlines() if l.strip()]
        heur = [{"skill": (line[:80] + "...") if len(line) > 80 else line, "description": line, "weight": 0.5, "min_years": None} for line in lines[:4]]
        if heur:
            return {"must_have": heur, "important": [], "nice_to_have": [], "implicit": []}
        return safe_empty

    try:
        raw_clean = html.unescape(raw).strip()
        raw_clean = re.sub(r"^```(?:json)?\n?|\n?```$", "", raw_clean).strip()
        first_brace = raw_clean.find("{")
        last_brace = raw_clean.rfind("}")
        if first_brace != -1 and last_brace != -1 and last_brace > first_brace:
            raw_json = raw_clean[first_brace:last_brace+1]
        else:
            raw_json = raw_clean
        parsed = json.loads(raw_json)
        return parsed
    except Exception:
        try:
            st.warning(f"{service} output couldn't be parsed as JSON.")
        except Exception:
            pass
        return safe_empty

def explain_candidate_with_llm(jd_struct: dict, resume_text: str, base_score: float, service: str) -> str:
    if not service: return "(LLM unavailable) Cannot generate explanation."

    system_prompt = (
        "You are an ATS explainability assistant. Given structured job requirements (JSON) and a candidate resume, "
        "produce a concise (<=200 words) explanation of why the candidate matches or doesn't, listing matched skills, missing skills, and an overall score 0-100. "
        "Crucially, **the last line of your output MUST be only the numeric match score (0-100)**, followed by no other text."
    )
    user_prompt = (
        f"JOB REQUIREMENTS JSON:\n{json.dumps(jd_struct, indent=2)}\n\nRESUME TEXT:\n{resume_text[:6000]}\n\`nBASE_SIM_SCORE:{base_score:.4f}"
    )
    raw = call_llm(service, system_prompt, user_prompt)
    
    if raw:
        return raw
    
    matched = []
    missing = []
    try:
        for group in ["must_have", "important", "nice_to_have"]:
            for item in jd_struct.get(group, []):
                sk = item.get("skill", "")
                if sk and sk.lower() in resume_text.lower():
                    matched.append(sk)
                else:
                    missing.append(sk)
    except Exception:
        pass
    matched = matched[:6]
    missing = missing[:6]
    
    fallback_llm_score = min(100, int(base_score * 100 * 1.5))
    
    return (
        f"({service} call failed) Heuristic summary — base_sim: {base_score:.3f}\n"
        f"Matched skills: {', '.join(matched) if matched else 'None'}\n"
        f"Missing skills (sample): {', '.join([m for m in missing if m]) if missing else 'None'}\n"
        f"{fallback_llm_score}"
    )


# ------------------- Streamlit UI -------------------

if 'resume_count' not in st.session_state:
    st.session_state.resume_count = 1

def add_resume_field():
    st.session_state.resume_count += 1

def remove_resume_field():
    if st.session_state.resume_count > 1:
        last_key = f"resume_text_area_{st.session_state.resume_count - 1}"
        if last_key in st.session_state:
            del st.session_state[last_key]
        st.session_state.resume_count -= 1


st.set_page_config(layout="wide", page_title="Candidate Matcher")
st.title("Candidate Matcher — Streamlit")

col1, col2 = st.columns([2, 1])

service_options = []
if os.getenv("GEMINI_API_KEY"):
    service_options.append("Gemini")
if os.getenv("OPENROUTER_API_KEY"):
    service_options.append("OpenRouter")

LLM_AVAILABLE = bool(service_options)


with col1:
    st.subheader("1. Job Description (JD)")
    jd_input_type = st.radio("JD input method:", ["Paste text", "Upload file"], index=0, key="jd_input_type")
    if jd_input_type == "Paste text":
        jd_text = st.text_area("Paste Job Description", height=300, key="jd_text_area")
    else:
        jd_file = st.file_uploader("Upload JD file (PDF/DOCX/TXT)", type=["pdf", "docx", "txt"], key="jd_file_uploader")
        jd_text = ""
        if jd_file:
            jd_text = extract_text_from_file(jd_file)

    st.markdown("---")

    # --- Resume Inputs ---
    st.subheader("2. Candidate Resumes")
    resume_input_type = st.radio("Resume input method:", ["Paste text", "Upload files"], index=0, key="resume_input_type")
        
    pasted_resumes = []
    uploaded_files = []

    if resume_input_type == "Paste text":
        st.markdown("**Paste Text Resumes:**")
        
        for i in range(st.session_state.resume_count):
            key = f"resume_text_area_{i}"
            if key not in st.session_state:
                st.session_state[key] = ""
        

        for i in range(st.session_state.resume_count):
            key = f"resume_text_area_{i}"
            
            st.text_area(f"Resume {i+1} Text", height=200, key=key)
            
            # Collect the valid text from session state
            if st.session_state[key].strip():
                pasted_resumes.append((f"Pasted Resume {i+1}", st.session_state[key]))


        col_add, col_remove = st.columns([1, 6])
        with col_add:
            st.button("➕ Add Resume", on_click=add_resume_field, key="add_resume_btn")
        with col_remove:
            if st.session_state.resume_count > 1:
                st.button("➖ Remove Last", on_click=remove_resume_field, key="remove_resume_btn")

    else: 
        uploaded_files = st.file_uploader("Upload candidate resumes (PDF/DOCX/TXT)", 
                                             type=["pdf", "docx", "txt"], 
                                             accept_multiple_files=True, 
                                             key="resume_file_uploader")

    st.markdown("---")
    run_btn = st.button("Run Matching")

with col2:
    st.markdown("**Settings**") 
    
    llm_service = None
    if LLM_AVAILABLE:

        if 'llm_service' not in st.session_state:
            st.session_state.llm_service = service_options[0]
            
        llm_service = st.radio("Select LLM Service:", service_options, key="llm_service")
        
        if llm_service == "OpenRouter":
            st.info(f"OpenRouter Model: {OPENROUTER_MODEL}")
        elif llm_service == "Gemini":
            st.info(f"Gemini Model: {GEMINI_MODEL}")
            
        use_llm_for_parsing = True
        use_llm_for_explanations = True
    else:
        use_llm_for_parsing = False
        use_llm_for_explanations = False
        st.warning("No LLM API keys configured. LLM features disabled.")
        
    st.markdown("---") 
    top_k = st.slider("Top K candidates to show explanations for", 1, 10, 5)


if run_btn:
    all_resumes_to_process = []
    
    # 1. Process uploaded files
    if uploaded_files:
        for f in uploaded_files:
            try:
                txt = extract_text_from_file(f)
                if txt.strip():
                    all_resumes_to_process.append((f.name, txt))
            except Exception as e:
                st.warning(f"Failed to process uploaded file {f.name}: {e}")

    # 2. Process pasted resumes
    if pasted_resumes:
        for name, text in pasted_resumes:
            if text.strip():
                all_resumes_to_process.append((name, text))
            
    
    if not jd_text or jd_text.strip() == "":
        st.error("Please provide a job description (paste text or upload file).")
    elif not all_resumes_to_process:
        st.error("Please provide at least one resume (paste text or upload file).")
    else:
        
        # --------------------- Main Processing Block (Loader) ---------------------
        with st.status("Processing and Matching Candidates...", expanded=True) as status:
            
            status.update(label="1. Resetting Data Store and Chroma...", state="running")
            # Reset store only if there are new candidates to process
            reset_local_store()
            reset_chroma_collection(force_delete_dir=False)

            status.update(label="2. Embedding and Storing Resumes...", state="running")
            # Add all collected resumes (files and pasted text)
            for name, txt in all_resumes_to_process:
                try:
                    add_resume_to_store(name, txt)
                except Exception as e:
                    st.warning(f"Failed to store resume {name}: {e}")

            status.update(label="3. Parsing Job Description Requirements...", state="running")
            if use_llm_for_parsing:
                try:
                    jd_struct = parse_jd_with_llm(jd_text, llm_service)
                except Exception as e:
                    st.error(f"{llm_service} call failed for JD parsing: {e}")
                    st.exception(traceback.format_exc())
                    jd_struct = {"must_have": [], "important": [], "nice_to_have": [], "implicit": []}
            else:
                jd_struct = {"must_have": [], "important": [], "nice_to_have": [], "implicit": []}
            
            status.update(label="4. Running Vector Similarity Search...", state="running")
            results = search_via_chroma_first_then_local(jd_text, top_k=max(top_k, len(all_resumes_to_process)))
            results = dedupe_results_by_id_keep_best(results)

            status.update(label="5. Preparing Final Scores...", state="running")
            
            final_candidates_data = []
            
            # Stage 1: Prepare for LLM
            for idx, r in enumerate(results):
                rid = r.get("id")
                score_vector = r.get("score", 0.0)
                payload = r.get("payload") or {}
                
                preview_text = _local_metadata[rid].get("preview", "") if rid in _local_metadata else r.get("preview", "")
                name = payload.get("name") if isinstance(payload, dict) else None
                name = name or (_local_metadata[rid]["name"] if rid in _local_metadata else f"resume_{rid}")
                
                candidate_data = {
                    "name": name, 
                    "text": preview_text, 
                    "score_vector": score_vector,
                    "score_llm_raw": 0.0,
                    "score_final": score_vector 
                }
                final_candidates_data.append(candidate_data)
            
            
            # Stage 2: LLM Explanation and Hybrid Scoring (Only top K candidates)
            if use_llm_for_explanations:
                status.update(label=f"6. Generating LLM Explanations and Hybrid Scores...", state="running")
                for i, c in enumerate(final_candidates_data[:top_k]):
                    try:
                        explanation_raw = explain_candidate_with_llm(jd_struct, c['text'], c['score_vector'], llm_service)

                        # 1. Extract LLM Score (last line of output)
                        parts = explanation_raw.strip().rsplit('\n', 1)
                        llm_score_text = parts[-1].strip()
                        explanation_display = explanation_raw if len(parts) == 1 else parts[0]
                        
                        llm_score_raw = 0.0
                        try:
                            llm_score_raw = float(re.sub(r'[^0-9.]', '', llm_score_text))
                            llm_score_normalized = min(1.0, max(0.0, llm_score_raw / 100.0))
                        except ValueError:
                            llm_score_normalized = 0.0
                            explanation_display = explanation_raw 
                        
                        # 2. Calculate Final Hybrid Score (50% Vector + 50% LLM)
                        score_vector = c['score_vector']
                        final_score = (0.5 * score_vector) + (0.5 * llm_score_normalized)

                        # Update data structures
                        c['score_llm_raw'] = llm_score_raw
                        c['score_final'] = final_score
                        c['explanation'] = explanation_display
                        
                    except Exception as e:
                        st.warning(f"Failed to get explanation for {c['name']}: {e}")
            
            status.update(label="7. Matching Complete! Displaying Results...", state="complete")

        # Stage 1: Display Vector Score (base similarity)
        st.subheader("Ranked candidates (by cosine similarity)")
        
        for idx, c in enumerate(final_candidates_data):
            st.markdown(f"**{idx+1}. {c['name']}** — base sim: {c['score_vector']:.4f}")
            with st.expander("Preview resume text"):
                st.write(c["text"][:2000] if len(c["text"]) > 2000 else c["text"])

        # Stage 2: LLM Explanation and Hybrid Scoring Display
        if use_llm_for_explanations:
            st.subheader(f"Top {min(top_k, len(final_candidates_data))} explanations (from {llm_service})")
            for i, c in enumerate(final_candidates_data[:top_k]):
                if c.get('score_llm_raw', 0) > 0:
                    # FIX: Display LLM Score separately in the header
                    st.markdown(f"**{c['name']}** — Final Hybrid Score: **{c['score_final']:.4f}** (LLM Score: **{c['score_llm_raw']:.1f}/100**)")
                    st.write(c['explanation'])
                else:
                    st.markdown(f"**{c['name']}** — Final Hybrid Score: **{c['score_final']:.4f}**")
                    st.write(c['explanation'])
        
        
        # Stage 3: Final Hybrid Score Table
        st.markdown("---")
        st.subheader("📊 Final Hybrid Matching Results")
        
        table_data = []
        for c in final_candidates_data:
            table_data.append({
                "Rank": "-",
                "Candidate Name": c['name'],
                "Vector Match (0.0-1.0)": f"{c['score_vector']:.4f}",
                "LLM Match (0-100)": f"{c['score_llm_raw']:.1f}" if c['score_llm_raw'] > 0 else "N/A",
                "Final Hybrid Score (0.0-1.0)": f"{c['score_final']:.4f}",
            })
            
        # Sort the table by Final Hybrid Score
        def get_sort_key(item):
            try:
                return float(item["Final Hybrid Score (0.0-1.0)"])
            except ValueError:
                return 0.0

        table_data_sorted = sorted(table_data, key=get_sort_key, reverse=True)

        # Apply the fix: Slice the sorted data to show only the Top K results specified by the user
        table_data_final_display = table_data_sorted[:top_k]

        for i, item in enumerate(table_data_final_display):
            item["Rank"] = i + 1
            
        st.dataframe(table_data_final_display, use_container_width=True)

        st.success("Analysis Complete!")

# Footer
st.markdown("---")
st.markdown("Notes: This application uses a Hybrid Matching system to accurately identify candidates whose resumes professionally align with the Job Description requirements.")